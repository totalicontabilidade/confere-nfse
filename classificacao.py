# -*- coding: utf-8 -*-
"""Classificação contábil das NFS-e tomadas a partir do plano de contas da empresa.

Para cada nota, escolhe:
  - a conta de DÉBITO (a despesa), pelo código de serviço da tabela nacional, pelo texto do
    serviço, pela discriminação e pelo nome do prestador;
  - a conta de CRÉDITO (o próprio fornecedor, que no Domínio tem conta reduzida própria);
  - as contas das RETENÇÕES (ISS, IRRF, PIS/COFINS/CSLL) a recolher.

Nada sai do computador: as regras casam conceitos com os NOMES das contas do plano, não com
códigos fixos de uma empresa. Por isso servem para o plano de qualquer cliente.
"""
import html
import re
import unicodedata

# ------------------------------------------------------------------ texto

def norm(s):
    """Maiúsculas, sem acento, só letras/números/espaço."""
    s = unicodedata.normalize("NFKD", str(s or "")).encode("ascii", "ignore").decode()
    s = re.sub(r"[^A-Za-z0-9]+", " ", s).upper()
    return re.sub(r"\s+", " ", s).strip()


SUFIXOS = {"LTDA", "ME", "EPP", "EIRELI", "SA", "S", "A", "SOCIEDADE", "INDIVIDUAL", "DE", "DA", "DO",
           "DOS", "DAS", "E", "EM", "RECUPERACAO", "JUDICIAL", "LIMITADA", "MEI", "CIA", "COMPANHIA"}


def nome_base(s):
    """Nome do fornecedor sem o que varia entre o XML e o cadastro: o CNPJ que o MEI põe na
    frente do nome ("47.124.487 DUILHO..."), o CPF no fim, e a forma jurídica."""
    t = norm(s)
    t = re.sub(r"^\d[\d ]{6,}\s+", "", t)       # 47 124 487 DUILHO -> DUILHO
    t = re.sub(r"\s+\d{11}$", "", t)            # ... 06187157540 -> ...
    return [p for p in t.split() if p not in SUFIXOS and not p.isdigit()]


# ------------------------------------------------------------------ plano de contas

RE_LINHA_DOMINIO = re.compile(r"^(\d+)\s+(S\s+)?(\d[\d.]*)\s+(.+?)\s+(\d)\s*$")


def ler_plano_pdf(caminho):
    """Relatório 'Plano de Contas' do Domínio: Código T Classificação Nome Grau."""
    import pdfplumber
    contas, empresa, cnpj = [], "", ""
    with pdfplumber.open(caminho) as pdf:
        for pg in pdf.pages:
            for linha in (pg.extract_text() or "").splitlines():
                l = linha.strip()
                if not empresa and l.upper().startswith("EMPRESA:"):
                    empresa = re.sub(r"(?i)^empresa:\s*|\s*folha.*$", "", l).strip()
                if not cnpj and re.match(r"(?i)^C\.?N\.?P\.?J", l):
                    cnpj = re.sub(r"\D", "", l)[:14]
                m = RE_LINHA_DOMINIO.match(l)
                if m:
                    contas.append({"cod": int(m[1]), "sintetica": bool(m[2]), "clas": m[3],
                                   "nome": m[4].strip(), "grau": int(m[5])})
    return {"empresa": empresa, "cnpj": cnpj, "contas": contas}


def ler_plano_tabela(linhas):
    """Planilha ou CSV: acha sozinho as colunas de código reduzido, classificação e nome."""
    linhas = [list(l) for l in linhas if l and any(c not in (None, "") for c in l)]
    cab_i, col = None, {}
    for i, l in enumerate(linhas[:15]):
        cab = [norm(c) for c in l]
        for j, c in enumerate(cab):
            if c in ("CODIGO", "COD", "REDUZIDO", "CODIGO REDUZIDO", "COD REDUZIDO", "CONTA"):
                col.setdefault("cod", j)
            elif "CLASSIFICA" in c:
                col.setdefault("clas", j)
            elif c in ("NOME", "DESCRICAO", "DESCRICAO DA CONTA", "NOME DA CONTA", "TITULO"):
                col.setdefault("nome", j)
            elif c in ("T", "TIPO"):
                col.setdefault("tipo", j)
        if {"cod", "clas", "nome"} <= col.keys():
            cab_i = i
            break
        col = {}
    if cab_i is None:
        raise ValueError("não achei as colunas de código, classificação e nome")
    contas = []
    for l in linhas[cab_i + 1:]:
        try:
            cod = int(str(l[col["cod"]]).strip().split(".")[0])
        except (ValueError, IndexError):
            continue
        clas = str(l[col["clas"]] or "").strip()
        nome = str(l[col["nome"]] or "").strip()
        if not clas or not nome:
            continue
        tipo = norm(l[col["tipo"]]) if "tipo" in col and col["tipo"] < len(l) else ""
        contas.append({"cod": cod, "sintetica": tipo.startswith("S"), "clas": clas, "nome": nome,
                       "grau": clas.count(".") + 1})
    return {"empresa": "", "contas": contas}


def ler_plano(caminho):
    c = caminho.lower()
    if c.endswith(".pdf"):
        return ler_plano_pdf(caminho)
    if c.endswith((".xlsx", ".xlsm")):
        import openpyxl
        wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
        return ler_plano_tabela(wb.worksheets[0].iter_rows(values_only=True))
    if c.endswith((".csv", ".txt")):
        import csv
        with open(caminho, encoding="utf-8-sig", errors="replace") as f:
            amostra = f.read(4096); f.seek(0)
            sep = ";" if amostra.count(";") > amostra.count(",") else ","
            return ler_plano_tabela(csv.reader(f, delimiter=sep))
    raise ValueError("formato de plano de contas não suportado: use o PDF do Domínio, Excel ou CSV")


def organizar(plano):
    """Separa o plano no que interessa: despesas, fornecedores e impostos a recolher."""
    contas = plano["contas"]
    por_clas = {c["clas"]: c for c in contas}

    def pai(c):
        partes = c["clas"].split(".")
        for k in range(len(partes) - 1, 0, -1):
            p = por_clas.get(".".join(partes[:k]))
            if p:
                return p
        return None

    def ancestrais(c):
        out, p = [], pai(c)
        while p:
            out.append(norm(p["nome"])); p = pai(p)
        return out

    analiticas = [c for c in contas if not c["sintetica"]]
    for c in analiticas:
        c["n"] = norm(c["nome"])
        c["_anc"] = ancestrais(c)
    despesas = [c for c in analiticas if any(a.startswith(("CUSTO", "DESPESA")) for a in c["_anc"])
                or c["clas"].split(".")[0] in ("3", "5")]
    fornecedores = [c for c in analiticas if any("FORNECEDOR" in a for a in c["_anc"])
                    and not any("ADIANTAMENTO" in a for a in c["_anc"])]
    for f in fornecedores:
        f["tokens"] = nome_base(f["nome"])
    tributos = [c for c in analiticas if any("OBRIGAC" in a and "TRIBUT" in a for a in c["_anc"])]
    return {"empresa": plano.get("empresa", ""), "cnpj": plano.get("cnpj", ""), "despesas": despesas, "fornecedores": fornecedores,
            "tributos": tributos, "todas": analiticas}


# ------------------------------------------------------------------ o que cada serviço é
# Cada conceito liga o que aparece na NOTA ao que aparece no NOME DA CONTA do plano.
#   conta:   regex sobre o nome normalizado da conta de despesa
#   texto:   regex sobre o texto da nota (serviço + discriminação + prestador)
#   itens:   itens da lista de serviços (LC 116), 4 dígitos, que apontam para o conceito
#   rotulo:  como a observação chama o serviço
CONCEITOS = [
    # forte: o que, aparecendo na DISCRIMINAÇÃO ou no nome do prestador, praticamente decide.
    # texto: indício de apoio. contas: nomes de conta aceitos, em ordem de preferência.
    dict(id="terceiros_isp", rotulo="Instalação / serviço técnico terceirizado",
         forte=r"INSTALAC\w*( E MANUTENCAO)?( DE)? (INTERNET|FIBRA|REDE)|RESPONSABILIDADE TECNICA|MESA DE SOM|COBERTURA FOTOGRAF|TRANSMISSAO D\w+ JOGOS|NARRACAO|INSTALACOES DE FIBRA|REINSTALAC",
         texto=r"INSTALAC|MAO DE OBRA|TERCEIRIZ", itens=[], contas=[r"PRESTADOS POR TERCEIROS", r"SERVICOS DE TERCEIROS"]),
    dict(id="saude", rotulo="Plano de saúde",
         forte=r"PLANOS? DE SAUDE|ASSISTENCIA MEDICA|HAPVIDA|UNIMED|COPARTICIPACAO|CONTRATO COLETIVO",
         texto=r"ODONTO|SAUDE", itens=["0422", "0423"], contas=[r"PLANOS? DE SAUDE", r"ASSISTENCIA MEDICA", r"DESPESAS MEDICAS"]),
    dict(id="advocacia", rotulo="Honorários advocatícios",
         forte=r"ADVOCA|ADVOGAD|HONORARIOS ADVOC", texto=r"JURIDIC", itens=["1714"], contas=[r"ADVOCAT"]),
    dict(id="contabil", rotulo="Honorários contábeis",
         forte=r"CONTABI|CONTADOR|ESCRITURACAO", texto=r"(?!)", itens=["1718", "1719"], contas=[r"CONTABE"]),
    dict(id="seg_trabalho", rotulo="Segurança do trabalho",
         forte=r"SEGURANCA DO TRABALHO|SESMT|PCMSO|\bPGR\b|LTCAT|MEDICINA DO TRABALHO|ADMISSIONA|METTASEG",
         texto=r"(?!)", itens=[], contas=[r"SEGURANCA DO TRABALHO"]),
    dict(id="vigilancia", rotulo="Vigilância e monitoramento",
         forte=r"VIGILANC|ALARME|SEGURANCA ELETRONICA|\bRONDA\b|MONITORAMENTO ELETRONICO|CFTV",
         texto=r"SEGURANCA|MONITORAMENTO", itens=["1102", "1103"], contas=[r"^DESPESAS COM SEGURANCA$", r"^SEGURANCA$", r"VIGILANC"]),
    dict(id="tv", rotulo="Canal / programação de TV",
         forte=r"\bCANAIS\b|\bCANAL\b|AUDIOVISU|SINAL E PROGRAMACAO|DISTRIBUICAO DE SINAL|PROGRAMADORA|\bCONTENT\b|RA TIM BUM|LICENCIAMENTO DE (PROGRAMACAO|CANA|CONTEUDO)|\w+TV\b|\bTV\b",
         texto=r"\bTV\b|TELEVISAO|PROGRAMACAO", itens=[], contas=[r"OPERADORAS? (DE )?TV", r"TELEVISAO"]),
    dict(id="processamento", rotulo="Processamento de dados",
         forte=r"PROCESSAMENTO|\bAPI\b|E ?SIM\b|GERENCIAMENTO DE COMUNICACAO DE DADOS",
         texto=r"(?!)", itens=["0103"], contas=[r"PROCESSAMENTO"]),
    dict(id="sistemas", rotulo="Sistema / licença de software",
         forte=r"MENSALIDADE|PLATAFORMA|LICENC(?!\w* DE (PROGRAMACAO|CANA|CONTEUDO|DISTRIBUICAO|SINAL))|SOFTWARE|\bSAAS\b|ASSINATURA|SUBSCRIPTION|\bSISTEMA \w+|SUPERAPP|\bIXC\w*|MODULO|DESENVOLVIMENTO DE SISTEMAS|\bSISTEMAS\b",
         texto=r"\bSISTEMAS?\b|APLICATIVO|\bAPP\b|FATURA MENSAL", itens=["0101", "0102", "0104", "0105"],
         contas=[r"SISTEMAS DE INFORMATICA", r"SOFTWARE", r"LICENC"]),
    dict(id="suporte", rotulo="Suporte técnico em informática",
         forte=r"^SUPORTE\b|SUPORTE -|ASSISTENCIA TECNICA|MANUTENCAO DE COMPUTADOR|MANUTENCAO DE MICRO|BANCO DE DADOS|FORMATAC|CONTRATO MENSAL|TROCA DE HD|\bSSD\b|WINDOWS|\bBIOS\b|MEMORIA RAM|MANUTENCAO BASICA",
         texto=r"SUPORTE|TECNOLOGIA DA INFORMACAO", itens=["0107", "1402"], contas=[r"SUPORTE TECNICO"]),
    dict(id="internet", rotulo="Internet / hospedagem / link",
         forte=r"HOSPEDAGEM DE SITE|HOSPEDAGEM|TRANSITO|IP TRANSIT|LINK DEDICADO|AREA DE TELECOMUNICAC|DATACENTER|DATA CENTER|\bISP\b|LOCAWEB|\bLWSA\b",
         texto=r"INTERNET|\bLINK\b|TELECOMUNICAC|BANDA LARGA|DOMINIO", itens=["0103", "0109"], contas=[r"INTERNET"]),
    dict(id="sva", rotulo="Serviço de valor adicionado",
         forte=r"\bSMS\b|\bSVA\b|VALOR ADICIONADO|VALOR AGREGADO", texto=r"(?!)", itens=[], contas=[r"VALOR AGREG"]),
    dict(id="publicidade", rotulo="Publicidade e propaganda",
         forte=r"PUBLICIDADE|PROPAGANDA|DIVULGA|ANUNCI|CARRO DE SOM|OUTDOOR|IMPULSIONA|BRANDPAGE|RECLAME AQUI|MARKETING|TRIEDRO|LETREIRO|PAINEL LUMINOSO|MIDIA EXTERIOR|\bSIGNS\b",
         texto=r"\bRADIO\b|MIDIA|CAMPANHA", itens=["1706", "1707", "1708", "1710"], contas=[r"PUBLICIDADE", r"PROPAGANDA", r"MARKETING"]),
    dict(id="patrocinio", rotulo="Patrocínio", forte=r"PATROCINIO", texto=r"(?!)", itens=[], contas=[r"PATROCINIO"]),
    dict(id="limpeza", rotulo="Higiene e limpeza",
         forte=r"LIMPEZA|HIGIENIZ|DEDETIZ|FAXINA", texto=r"(?!)", itens=["0710", "0713"], contas=[r"LIMPEZA"]),
    dict(id="tarifas", rotulo="Tarifa bancária / de pagamento",
         forte=r"TARIFA|INSTITUICAO DE PAGAMENTO", texto=r"COBRANCA", itens=[f"15{i:02d}" for i in range(1, 19)],
         contas=[r"TARIFAS BANCARIAS", r"TARIFA"]),
    dict(id="boletos", rotulo="Tarifa de boleto",
         forte=r"TARIFA\w* (S |SOBRE |DE )?BOLETO|EMISSAO DE BOLETO|REGISTRO DE BOLETO", texto=r"(?!)", itens=[], contas=[r"BOLETO"]),
    dict(id="postes", rotulo="Aluguel de postes",
         forte=r"POSTES?\b|COMPARTILHAMENTO DE INFRAESTRUTURA|PONTO DE FIXACAO|OCUPACAO DE INFRA", texto=r"(?!)", itens=[], contas=[r"POSTES"]),
    dict(id="monit_redes", rotulo="Monitoramento de redes",
         forte=r"MONITORAMENTO DE REDE|REPARO DE REDE|REPARO DE URGENCIA|\bNOC\b|ROMPIMENTO", texto=r"(?!)", itens=[], contas=[r"MONITORAMENTO DE REDES"]),
    dict(id="teleatendimento", rotulo="Teleatendimento",
         forte=r"CALL ?CENTER|TELEATENDIMENTO|CENTRAL DE ATENDIMENTO", texto=r"\bSAC\b", itens=[], contas=[r"TELEATENDIMENTO", r"CALL ?CENTER"]),
    dict(id="treinamento", rotulo="Curso / treinamento",
         forte=r"\bCURSO|TREINAMENTO|CAPACITAC|PALESTRA|INSCRICAO|WORKSHOP", texto=r"(?!)", itens=["0801", "0802"], contas=[r"CURSOS", r"TREINAMENTO"]),
    dict(id="viagem", rotulo="Viagem / estadia",
         forte=r"HOTEL|DIARIAS?\b|PASSAGE|POUSADA|HOSPEDAGEM EM", texto=r"TURISMO|VIAGE", itens=["0901", "0902"], contas=[r"VIAGE", r"ESTADIA"]),
    dict(id="man_veiculo", rotulo="Manutenção de veículo",
         forte=r"VEICUL|AUTOMOTIV|OFICINA|MECANIC|\bCARRO\b|\bPLACA\b|\bCAMBIO\b|FUNILARIA|\bPNEUS?\b|LAVAGEM|\bMOTOS?\b|STRADA|OROCH",
         texto=r"TRANSMISSO", itens=[], contas=[r"MANUTENCAO DE VEICULO"]),
    dict(id="man_maquina", rotulo="Manutenção de máquinas e equipamentos",
         forte=r"NOBREAK|GERADOR|AR CONDICIONADO|REFRIGERAC|MANUTENCAO DE (MAQUINA|EQUIPAMENTO)|IMPRESSORA|MULTIFUNCIONAL|TONER|CABECA DE IMPRESSAO|EPSON|BROTHER",
         texto=r"EQUIPAMENTO|MAQUINA", itens=["1401", "1402", "1406"], contas=[r"MANUTENCAO DE MAQUINAS"]),
    dict(id="reforma", rotulo="Manutenção predial / reforma",
         forte=r"REFORMA (DO|DA|DE) (PREDIO|LOJA|IMOVEL|SALA)|PINTURA|INSTALACAO ELETRICA|HIDRAULIC|ALVENARIA|PELICU",
         texto=r"\bOBRA\b|CONSTRUC", itens=["0702", "0705"], contas=[r"MANUTENCAO E REFORMA", r"REFORMA"]),
    dict(id="graficos", rotulo="Impressos gráficos",
         forte=r"GRAFIC|PANFLETO|BANNER|ADESIVO|COMUNICACAO VISUAL|PLOTAGEM|\bIMPRESSOS\b", texto=r"\bIMPRESSOS?\b", itens=["1305", "2402"],
         contas=[r"IMPRESSOS", r"GRAFIC"]),
    dict(id="loc_equip", rotulo="Locação de equipamentos", forte=r"(LOCACAO|ALUGUEL) DE (EQUIPAMENTO|MAQUINA)",
         texto=r"(?!)", itens=["0301"], contas=[r"(ALUGUEL|LOCACAO) DE EQUIPAMENTO"]),
    dict(id="loc_veiculo", rotulo="Locação de veículo", forte=r"LOCACAO DE VEICULO|LOCADORA", texto=r"(?!)", itens=[], contas=[r"LOCACAO DE VEICULO"]),
    dict(id="associacao", rotulo="Sindicato / associação", forte=r"SINDICATO|ANUIDADE|CONTRIBUICAO ASSOCIATIVA",
         texto=r"ASSOCIACAO", itens=[], contas=[r"SINDICATO", r"ASSOCIAC", r"ENTIDADE"]),
    dict(id="seguros", rotulo="Seguro", forte=r"\bSEGURO\b|APOLICE", texto=r"(?!)", itens=[], contas=[r"^SEGUROS?$", r"DESPESA COM SEGUROS"]),
    dict(id="telefone", rotulo="Telefonia", forte=r"TELEFONIA|\bVOIP\b|LINHA FIXA", texto=r"TELEFON", itens=[], contas=[r"^TELEFONE"]),
    dict(id="consultoria", rotulo="Consultoria",
         forte=r"CONSULTORIA|ASSESSORIA", texto=r"(?!)", itens=["1701", "0106"], contas=[r"CONSULTORIA"]),
    dict(id="exterior", rotulo="Serviço do exterior", forte=r"INVOICE", texto=r"EXTERIOR", itens=[], contas=[r"EXTERIOR", r"INVOICE"]),
    # genéricos: só vencem quando nada específico casou
    dict(id="terceiros", rotulo="Serviço de terceiros", forte=r"(?!)", texto=r"SERVICO|PRESTACAO", itens=[],
         contas=[r"PRESTADOS POR TERCEIROS", r"SERVICOS DE TERCEIROS"], generico=True),
    dict(id="diversas", rotulo="Despesa diversa", forte=r"(?!)", texto=r"(?!)", itens=[],
         contas=[r"DESPESAS DIVERSAS", r"OUTRAS DESPESAS"], generico=True),
]
for _c in CONCEITOS:
    _c["_forte"] = re.compile(_c["forte"])
    _c["_texto"] = re.compile(_c["texto"])
    _c["_contas"] = [re.compile(x) for x in _c["contas"]]


def item_lc116(cod):
    """'130301' (cTribNac) ou '13.03' -> '1303'."""
    d = re.sub(r"\D", "", str(cod or ""))
    return d[:4] if len(d) >= 4 else ""


def conta_do_conceito(conceito, despesas):
    """A primeira conta do plano que casa, respeitando a ordem de preferência do conceito."""
    for rx in conceito["_contas"]:
        for c in despesas:
            if rx.search(c["n"]):
                return c
    return None


def limpar_discriminacao(discr, servico):
    """Muito prestador cola na discriminação o próprio texto do código de serviço ("Suporte
    técnico, manutenção e outros serviços em tecnologia da informação"). Isso é o código de
    novo, não informação — então sai, para não valer duas vezes."""
    d = norm(discr)
    for trecho in re.split(r"[,;.]", str(servico or "")):
        t = norm(trecho)
        if len(t) >= 18 and t in d:
            d = d.replace(t, " ")
    # o que é instrução de pagamento ou aviso legal, não descrição do serviço
    ruido = (r"VALOR APROX\w* (DOS )?TRIBUTOS.*|LEI 12 741.*|DADOS BANCARIOS.*|AGENCIA \d.*|CONTA CORRENTE.*"
             r"|PAGAMENTO VIA PIX.*|CHAVE (PIX|CNPJ).*|\bPIX \d.*|VENCIMENTO \d.*|VENCTO \d.*|BANCO \d{3}.*")
    return re.sub(ruido, " ", d)


# Pesos: a discriminação é o que o prestador escreveu sobre ESTE serviço; o código e a
# descrição do código são escolhidos de qualquer jeito e valem pouco.
P_DISCR_FORTE, P_DISCR, P_NOME_FORTE, P_NOME, P_ITEM, P_SERVICO = 4.0, 1.5, 3.0, 1.0, 1.5, 0.75


def pontuar(nota, despesas):
    """Devolve [(pontos, conceito, conta, sinais, decisivo)] do melhor para o pior."""
    servico = norm(nota.get("servico"))
    discr = limpar_discriminacao(nota.get("discriminacao"), nota.get("servico"))
    prest = norm(nota.get("prestador"))
    item = item_lc116(nota.get("cod_servico"))
    out = []
    for c in CONCEITOS:
        conta = conta_do_conceito(c, despesas)
        if not conta:
            continue                                   # o plano desta empresa não tem essa conta
        sinais, pts, decisivo = [], 0.0, False
        if c["_forte"].search(discr):
            pts += P_DISCR_FORTE; sinais.append("discriminação"); decisivo = True
        elif c["_texto"].search(discr):
            pts += P_DISCR; sinais.append("discriminação")
        if c["_forte"].search(prest):
            pts += P_NOME_FORTE; sinais.append("nome do prestador"); decisivo = True
        elif c["_texto"].search(prest):
            pts += P_NOME; sinais.append("nome do prestador")
        if item and item in c["itens"]:
            pts += P_ITEM; sinais.append(f"código {item[:2]}.{item[2:]}")
        if c["_forte"].search(servico) or c["_texto"].search(servico):
            pts += P_SERVICO
        if c.get("generico"):
            pts = pts * 0.3 + 0.2
        if pts <= 0:
            continue
        out.append((pts, c, conta, sinais, decisivo))
    out.sort(key=lambda t: -t[0])
    return out


# ------------------------------------------------------------------ fornecedor (crédito)

COMUNS = {"SERVICOS", "SERVICO", "COMERCIO", "COMERCIAL", "AUTO", "CENTRAL", "CASA", "BANCO", "GRUPO",
          "INSTITUTO", "EMPRESA", "BRASIL", "NACIONAL", "SISTEMAS", "TECNOLOGIA", "DISTRIBUIDORA", "INDUSTRIA",
          "POSTO", "SUPERMERCADO", "FARMACIA", "LOJA", "CLINICA", "HOSPITAL", "ASSOCIACAO", "FUNDACAO"}


def _mesma_palavra(a, b):
    """'FUND' e 'FUNDACAO', 'PAULIS' e 'PAULISTA': o cadastro abrevia."""
    if a == b:
        return True
    curto, longo = (a, b) if len(a) < len(b) else (b, a)
    return len(curto) >= 4 and longo.startswith(curto)


def _casadas(xs, ys):
    usados, n = set(), 0
    for x in xs:
        for j, y in enumerate(ys):
            if j not in usados and _mesma_palavra(x, y):
                usados.add(j); n += 1
                break
    return n


def achar_fornecedor(nome, fornecedores):
    alvo = nome_base(nome)
    if not alvo:
        return None, 0.0
    melhor, nota = None, 0.0
    for f in fornecedores:
        tf = f["tokens"]
        if not tf:
            continue
        inter = _casadas(alvo, tf)
        if not inter:
            continue
        jac = inter / (len(alvo) + len(tf) - inter)
        cont = inter / min(len(alvo), len(tf))           # um nome contido no outro
        sc = 0.55 * jac + 0.45 * cont
        if _mesma_palavra(alvo[0], tf[0]):
            sc += 0.15                                   # a primeira palavra do nome pesa mais
            if alvo[0] not in COMUNS and len(alvo[0]) >= 5:
                sc += 0.12                               # e mais ainda se for rara (ZENVIA, MXCSOFT)
        if sc > nota:
            melhor, nota = f, sc
    return (melhor, min(1.0, nota)) if nota >= 0.55 else (None, nota)


# ------------------------------------------------------------------ retenções (crédito)

def contas_de_retencao(tributos):
    def acha(*padroes):
        for p in padroes:
            r = re.compile(p)
            for c in tributos:
                if r.search(c["n"]):
                    return c
        return None
    return {
        "iss": acha(r"^ISS\b.*RECOLHER", r"\bISS\b.*RETID", r"\bISS\b"),
        "irrf": acha(r"^IRRF\b.*RECOLHER", r"IMPOSTO DE RENDA RETIDO", r"\bIRRF\b"),
        "csrf": acha(r"CONTRIBU\w* SOCIAIS RETIDAS", r"PIS.*COFINS.*CSLL", r"\bCSRF\b", r"CONTRIBU\w* RETIDAS"),
    }


# ------------------------------------------------------------------ a nota inteira

def classificar(nota, org, memoria=None):
    """nota: dict com cnpj, prestador, cod_servico, servico, discriminacao, valor, iss_retido,
    irrf_retido, csrf_retido, cancelada, substituida."""
    r = {"debito": None, "credito": None, "confianca": "Baixa", "observacao": "", "alternativa": None}
    obs = []

    if nota.get("cancelada") or nota.get("substituida"):
        r["confianca"] = "Excluir"
        obs.append("Nota cancelada — não lançar." if nota.get("cancelada") else "Nota substituída por outra — lançar só a substituta.")

    # 1) o que já foi decidido para este fornecedor vale mais que qualquer regra
    hist = (memoria or {}).get(re.sub(r"\D", "", str(nota.get("cnpj") or ""))) or {}
    if isinstance(hist, int):
        hist = {str(hist): 1}
    usadas = sorted(((int(k), n) for k, n in hist.items()), key=lambda kv: -kv[1])
    por_cod = {c["cod"]: c for c in org["despesas"]}
    usadas = [(por_cod[k], n) for k, n in usadas if k in por_cod]
    ranking = pontuar(nota, org["despesas"])
    if usadas:
        conta = usadas[0][0]
        if conta:
            r["debito"] = conta
            if len(usadas) == 1:
                if r["confianca"] != "Excluir":
                    r["confianca"] = "Alta"
                obs.append("Mesma conta já usada para este fornecedor.")
            else:
                # o fornecedor presta mais de um serviço: a nota decide, e você confere
                texto = next((c for c, _ in usadas if ranking and ranking[0][2]["cod"] == c["cod"]), None)
                if texto and texto["cod"] != conta["cod"]:
                    conta = texto; r["debito"] = conta
                outra = next(c for c, _ in usadas if c["cod"] != conta["cod"])
                r["alternativa"] = outra
                if r["confianca"] != "Excluir":
                    r["confianca"] = "Média"
                obs.append(f"Este fornecedor já foi lançado em {len(usadas)} contas diferentes; "
                           f"pela nota, {conta['nome'].lower()}. Alternativa: {outra['cod']} {outra['nome']}.")
            # se a própria nota aponta com clareza para outra conta, vale conferir
            forte = next((t for t in ranking if t[4] and not t[1].get("generico")), None)
            if forte and forte[2]["cod"] != conta["cod"] and forte is ranking[0]:
                r["alternativa"] = forte[2]
                if r["confianca"] == "Alta":
                    r["confianca"] = "Média"
                obs.append(f"Mas esta nota parece {forte[1]['rotulo'].lower()} ({forte[2]['cod']} {forte[2]['nome']}) — conferir.")

    if not r["debito"]:
        if ranking:
            pts, conc, conta, sinais, decisivo = ranking[0]
            r["debito"] = conta
            seg = next((t for t in ranking[1:] if t[2]["cod"] != conta["cod"]), None)
            margem = pts - (seg[0] if seg else 0)
            if r["confianca"] != "Excluir":
                if conc.get("generico"):
                    r["confianca"] = "Baixa"
                elif decisivo:
                    # Regra nunca dá Alta: medido em outra empresa, só 59% das Altas por regra
                    # acertavam. Alta é da memória — o que você já confirmou para este fornecedor.
                    r["confianca"] = "Média"
                else:
                    r["confianca"] = "Baixa"
            obs.append(f"{conc['rotulo']} — pela {' e '.join(sinais)}." if sinais else f"{conc['rotulo']}.")
            r["fornecedor_novo"] = True
            if seg and not seg[1].get("generico") and (margem < 2.0 or seg[4]):
                r["alternativa"] = seg[2]
                obs.append(f"Alternativa: {seg[2]['cod']} {seg[2]['nome']}.")
        else:
            r["debito"] = next((c for c in org["despesas"] if re.search(r"DESPESAS DIVERSAS", c["n"])), None)
            obs.append("Não reconheci o serviço — conferir a conta.")

    # 2) fornecedor
    forn, sim = achar_fornecedor(nota.get("prestador"), org["fornecedores"])
    r["credito"] = forn
    r["sim_fornecedor"] = sim
    if not forn:
        obs.append("Fornecedor não encontrado no plano — cadastrar no Domínio.")
        if r["confianca"] == "Alta":
            r["confianca"] = "Média"
    elif sim < 0.8:
        obs.append(f"Fornecedor casado pelo nome ({forn['nome']}) — confirmar.")

    # 3) retenções
    ret = contas_de_retencao(org["tributos"])
    partes = []
    if (nota.get("iss_retido") or 0) > 0 and ret["iss"]:
        partes.append(f"ISS: {ret['iss']['cod']} {ret['iss']['nome']}")
    if (nota.get("irrf_retido") or 0) > 0 and ret["irrf"]:
        partes.append(f"IRRF: {ret['irrf']['cod']} {ret['irrf']['nome']}")
    if (nota.get("csrf_retido") or 0) > 0 and ret["csrf"]:
        partes.append(f"CSRF: {ret['csrf']['cod']} {ret['csrf']['nome']}")
    r["retencoes"] = " | ".join(partes)

    r["observacao"] = " ".join(obs)
    return r


# ------------------------------------------------------------------ as notas, a partir dos XMLs

def _tag(xml, *nomes):
    for n in nomes:
        m = re.search(rf"<(?:\w+:)?{n}>([^<]*)</(?:\w+:)?{n}>", xml)
        if m and m.group(1).strip():
            return html.unescape(m.group(1).strip())      # &amp; -> &, &quot; -> "
    return ""


def _bloco(xml, *nomes):
    for n in nomes:
        m = re.search(rf"<(?:\w+:)?{n}\b[^>]*>(.*?)</(?:\w+:)?{n}>", xml, re.S)
        if m:
            return m.group(1)
    return ""


def _num(v):
    try:
        return round(float(str(v).replace(",", ".")), 2)
    except (TypeError, ValueError):
        return 0.0


def nota_do_xml(xml, arquivo=""):
    """Uma NFS-e (padrão nacional ou ABRASF) no formato que o classificador usa. None se não for nota."""
    if re.search(r"<(?:\w+:)?evento\b", xml[:600]):
        return None
    nacional = "<infNFSe" in xml or "<nNFSe>" in xml
    if nacional:
        emit = _bloco(xml, "emit") or _bloco(xml, "prest")
        chave = re.sub(r"\D", "", (re.search(r'<infNFSe[^>]*Id="([^"]+)"', xml) or [None, ""])[1])
        ret_iss = _tag(xml, "tpRetISSQN")
        toma = _bloco(xml, "toma")
        n = dict(
            numero=_tag(xml, "nNFSe"), chave=chave if len(chave) == 50 else "",
            competencia=_tag(xml, "dCompet")[:10],
            cnpj=_tag(emit, "CNPJ", "CPF"), prestador=_tag(emit, "xNome"),
            cod_servico=_tag(xml, "cTribNac"), servico=_tag(xml, "xTribNac"),
            discriminacao=_tag(xml, "xDescServ"),
            valor=_num(_tag(xml, "vServ")),
            iss_retido=_num(_tag(xml, "vISSQN")) if ret_iss in ("2", "3") else 0.0,
            irrf_retido=_num(_tag(xml, "vRetIRRF")),
            csrf_retido=_num(_tag(xml, "vRetCSLL")),   # no nacional, vRetCSLL já soma PIS+COFINS+CSLL
            substitui=re.sub(r"\D", "", _tag(xml, "chSubstda")),
            tomador=re.sub(r"\D", "", _tag(toma, "CNPJ", "CPF")), tomador_nome=_tag(toma, "xNome"),
            cancelada=False,
        )
    else:
        prest = _bloco(xml, "PrestadorServico", "Prestador")
        iss_ret = _tag(xml, "IssRetido") == "1"
        pis, cof, csll = (_num(_tag(xml, t)) for t in ("ValorPis", "ValorCofins", "ValorCsll"))
        n = dict(
            numero=_tag(xml, "Numero"), chave="",
            competencia=(_tag(xml, "Competencia") or _tag(xml, "DataEmissao"))[:10],
            cnpj=_tag(prest, "Cnpj", "Cpf"), prestador=_tag(prest, "RazaoSocial", "NomeFantasia"),
            cod_servico=_tag(xml, "ItemListaServico", "CodigoTributacaoMunicipio"), servico="",
            discriminacao=_tag(xml, "Discriminacao"),
            valor=_num(_tag(xml, "ValorServicos")),
            iss_retido=_num(_tag(xml, "ValorIssRetido") or (_tag(xml, "ValorIss") if iss_ret else 0)),
            irrf_retido=_num(_tag(xml, "ValorIr")),
            csrf_retido=round(pis + cof + csll, 2),
            tomador=re.sub(r"\D", "", _tag(_bloco(xml, "TomadorServico", "Tomador"), "Cnpj", "Cpf")),
            tomador_nome=_tag(_bloco(xml, "TomadorServico", "Tomador"), "RazaoSocial"),
            substitui="", cancelada=bool(re.search(r"<(?:\w+:)?(NfseCancelamento|Cancelamento)\b", xml)),
        )
    n["arquivo"] = arquivo
    n["cnpj"] = re.sub(r"\D", "", n["cnpj"])
    return n if n["numero"] or n["valor"] else None


def notas_dos_xmls(textos):
    """[(nome, texto)] -> notas, já marcando canceladas (evento do portal) e substituídas."""
    notas, canceladas = [], set()
    for nome, xml in textos:
        if re.search(r"<(?:\w+:)?evento\b", xml[:600]):
            ch = re.sub(r"\D", "", _tag(xml, "chNFSe"))
            motivo = norm(_tag(xml, "xMotivo") + " " + _tag(xml, "xDesc"))
            if ch and ("e101101" in xml or "CANCEL" in motivo):
                canceladas.add(ch)
            continue
        n = nota_do_xml(xml, nome)
        if n:
            notas.append(n)
    substituidas = {n["substitui"] for n in notas if n["substitui"]}
    vistos = set()
    unicas = []
    for n in notas:
        chave = n["chave"] or f"{n['cnpj']}|{n['numero']}"
        if chave in vistos:
            continue                                   # o mesmo XML carregado duas vezes
        vistos.add(chave)
        if n["chave"] and n["chave"] in canceladas:
            n["cancelada"] = True
        if n["chave"] and n["chave"] in substituidas:
            n["substituida"] = True
        unicas.append(n)
    unicas.sort(key=lambda n: (n["competencia"], n["prestador"]))
    return unicas


# ------------------------------------------------------------------ memória: aprender com o que já foi classificado

def ler_classificacao_anterior(caminho, org):
    """Uma planilha já conferida pelo contador vira memória: CNPJ do fornecedor -> conta de débito.
    Aceita o modelo deste sistema (com a coluna do código reduzido) e planilhas em que a conta
    aparece como 'classificação nome' (3.2.10.002.078 CAMPEONATO SERGIPANO)."""
    import openpyxl
    wb = openpyxl.load_workbook(caminho, read_only=True, data_only=True)
    por_cod = {c["cod"]: c for c in org["despesas"]}
    por_clas = {c["clas"]: c for c in org["despesas"]}
    por_nome = {c["n"]: c for c in org["despesas"]}
    votos = {}

    def conta_de(v):
        if v is None:
            return None
        t = str(v).strip()
        if re.fullmatch(r"\d+(\.0)?", t):
            return por_cod.get(int(float(t)))
        m = re.match(r"^(\d[\d.]*\d)\b\s*(.*)$", t)
        if m and m[1] in por_clas:
            return por_clas[m[1]]
        return por_nome.get(norm(re.sub(r"^\d+\s+", "", t)))

    for ws in wb.worksheets:
        linhas = list(ws.iter_rows(values_only=True))
        for i, l in enumerate(linhas[:12]):
            cab = [norm(c) for c in (l or [])]
            c_doc = next((j for j, c in enumerate(cab) if "CNPJ" in c or c in ("CPF", "DOCUMENTO")), None)
            c_cod = next((j for j, c in enumerate(cab) if "DEBITO" in c and ("REDUZ" in c or "COD" in c)), None)
            c_conta = next((j for j, c in enumerate(cab) if ("CONTA" in c and ("CONTABIL" in c or "DESPESA" in c))
                            or c == "CONTA CONTABIL"), None)
            c_conf = next((j for j, c in enumerate(cab) if c.startswith("CONFIANCA")), None)
            alvo = c_cod if c_cod is not None else c_conta
            if c_doc is None or alvo is None:
                continue
            for l2 in linhas[i + 1:]:
                if not l2 or c_doc >= len(l2) or alvo >= len(l2):
                    continue
                doc = re.sub(r"\D", "", str(l2[c_doc] or ""))
                if len(doc) not in (11, 14):
                    continue
                if c_conf is not None and c_conf < len(l2) and norm(l2[c_conf]) == "EXCLUIR":
                    continue
                conta = conta_de(l2[alvo])
                if conta:
                    votos.setdefault(doc, {}).setdefault(conta["cod"], 0)
                    votos[doc][conta["cod"]] += 1
            break
    wb.close()
    # guarda todas as contas que o fornecedor já usou, com quantas vezes
    return {doc: {str(k): n for k, n in v.items()} for doc, v in votos.items()}


def juntar_memoria(antiga, nova):
    """Soma o que já se sabia com o que acabou de ser aprendido."""
    out = {d: dict(v) for d, v in (antiga or {}).items()}
    for d, v in (nova or {}).items():
        alvo = out.setdefault(d, {})
        for k, n in v.items():
            alvo[str(k)] = alvo.get(str(k), 0) + n
    return out


# ------------------------------------------------------------------ o Excel, no modelo da Totali

def gerar_excel(notas, org, empresa="", competencia="", memoria=None):
    import openpyxl
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter

    AZUL = "FF1F3864"
    CORES = {"Alta": "FFE2EFDA", "Média": "FFFFF2CC", "Baixa": "FFFCE4D6", "Excluir": "FFD9D9D9"}
    MOEDA = "#,##0.00;[RED]\\-#,##0.00;\\-"
    fina = Side(style="thin", color="FFBFBFBF")
    borda = Border(left=fina, right=fina, top=fina, bottom=fina)
    cab_fonte = Font(name="Arial", size=10, bold=True, color="FFFFFFFF")
    cab_fundo = PatternFill("solid", fgColor=AZUL)
    titulo = Font(name="Arial", size=14, bold=True, color=AZUL)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Classificação NFS-e"
    nome_emp = empresa or org.get("empresa") or "Empresa"
    ws["A1"] = f"{nome_emp} — Classificação contábil das NFS-e recebidas" + (f" (competência {competencia})" if competencia else "")
    ws["A1"].font = titulo
    conhecidos = sum(1 for n in notas if (memoria or {}).get(n["cnpj"]))
    ws["A2"] = (f"Base: {len(notas)} XMLs de NFS-e x Plano de Contas {nome_emp}. "
                f"Lançamento sugerido: D despesa / C fornecedor; retenções a crédito nas contas de impostos retidos. "
                f"{conhecidos} nota(s) de fornecedor já classificado antes; as demais são sugestão das regras.")
    ws["A2"].font = Font(name="Arial", size=9)

    cab = ["Nº NFS-e", "Competência", "CNPJ/CPF prestador", "Prestador", "Cód. serviço", "Serviço (tabela nacional)",
           "Discriminação da nota", "Valor do serviço", "ISS retido", "IRRF retido", "PIS/COFINS/CSLL retidos",
           "Valor líquido a pagar", "Débito — cód. reduzido", "Débito — classificação", "Débito — conta",
           "Crédito — cód. reduzido (fornecedor)", "Crédito — conta fornecedor", "Confiança", "Observação / o que validar"]
    larg = [15, 11, 16, 32, 9, 30, 55, 13, 10, 10, 12, 13, 10, 17, 30, 12, 32, 10, 60]
    for j, (t, w) in enumerate(zip(cab, larg), start=1):
        c = ws.cell(4, j, t)
        c.font, c.fill = cab_fonte, cab_fundo
        c.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
        ws.column_dimensions[get_column_letter(j)].width = w
    ws.row_dimensions[4].height = 42

    fonte = Font(name="Arial", size=9)
    resultados = []
    linha = 5
    for n in notas:
        r = classificar(n, org, memoria)
        resultados.append((linha, n, r))
        d, c = r["debito"], r["credito"]
        valores = [n["numero"], n["competencia"], n["cnpj"], n["prestador"], n["cod_servico"], n["servico"],
                   n["discriminacao"], n["valor"], n["iss_retido"], n["irrf_retido"], n["csrf_retido"],
                   f"=H{linha}-I{linha}-J{linha}-K{linha}",
                   d["cod"] if d else None, d["clas"] if d else None, d["nome"] if d else None,
                   c["cod"] if c else None, c["nome"] if c else None, r["confianca"], r["observacao"]]
        for j, v in enumerate(valores, start=1):
            cel = ws.cell(linha, j, v)
            cel.font = fonte
            cel.border = borda
            cel.alignment = Alignment(vertical="top", wrap_text=(j in (7, 19)))
            if 8 <= j <= 12:
                cel.number_format = MOEDA
        conf = ws.cell(linha, 18)
        conf.fill = PatternFill("solid", fgColor=CORES.get(r["confianca"], "FFFFFFFF"))
        if r["confianca"] == "Excluir":
            conf.font = Font(name="Arial", size=9, color="FF808080")
        linha += 1
    ultima = linha - 1

    tot = ultima + 1
    ws.cell(tot, 7, "TOTAL (sem as notas a excluir)").font = Font(name="Arial", size=11, bold=True)
    for j in range(8, 13):
        col = get_column_letter(j)
        cel = ws.cell(tot, j, f'=SUMIFS({col}5:{col}{ultima},$R$5:$R${ultima},"<>Excluir")')
        cel.font = Font(name="Arial", size=11, bold=True)
        cel.number_format = "#,##0.00"
    ws.freeze_panes = "E5"
    ws.auto_filter.ref = f"A4:S{ultima}"

    # ---------------- resumo por conta de despesa
    rs = wb.create_sheet("Resumo por conta")
    rs["A1"] = f"Resumo por conta de despesa — NFS-e {competencia}".strip(" —")
    rs["A1"].font = titulo
    for j, t in enumerate(["Cód. reduzido", "Classificação", "Conta", "Qtd. notas", "Valor do serviço", "Valor líquido"], 1):
        cel = rs.cell(3, j, t)
        cel.font = Font(name="Arial", size=11, bold=True, color="FFFFFFFF"); cel.fill = cab_fundo
        cel.alignment = Alignment(wrap_text=True, horizontal="center")
    for col, w in zip("ABCDEF", (12, 18, 48, 10, 16, 16)):
        rs.column_dimensions[col].width = w
    usadas = sorted({(r["debito"]["clas"], r["debito"]["cod"], r["debito"]["nome"])
                     for _, _, r in resultados if r["debito"] and r["confianca"] != "Excluir"})
    aba = "'Classificação NFS-e'"
    lr = 4
    for clas, cod, nome in usadas:
        rs.cell(lr, 1, cod); rs.cell(lr, 2, clas); rs.cell(lr, 3, nome)
        rs.cell(lr, 4, f'=COUNTIFS({aba}!$M$5:$M${ultima},A{lr},{aba}!$R$5:$R${ultima},"<>Excluir")')
        rs.cell(lr, 5, f'=SUMIFS({aba}!$H$5:$H${ultima},{aba}!$M$5:$M${ultima},A{lr},{aba}!$R$5:$R${ultima},"<>Excluir")')
        rs.cell(lr, 6, f'=SUMIFS({aba}!$L$5:$L${ultima},{aba}!$M$5:$M${ultima},A{lr},{aba}!$R$5:$R${ultima},"<>Excluir")')
        for j in range(1, 7):
            rs.cell(lr, j).font = Font(name="Arial", size=10)
        for j in (5, 6):
            rs.cell(lr, j).number_format = "#,##0.00"
        lr += 1
    rs.cell(lr, 3, "TOTAL").font = Font(name="Arial", size=10, bold=True)
    for j, col in ((4, "D"), (5, "E"), (6, "F")):
        cel = rs.cell(lr, j, f"=SUM({col}4:{col}{lr - 1})")
        cel.font = Font(name="Arial", size=10, bold=True)
        if j > 4:
            cel.number_format = "#,##0.00"
    rs.cell(lr + 1, 3, "Conferência (deve ser zero)").font = Font(name="Arial", size=10, italic=True)
    cel = rs.cell(lr + 1, 5, f"=E{lr}-{aba}!H{tot}")
    cel.number_format = "#,##0.00"

    # ---------------- retenções
    rt = wb.create_sheet("Retenções a recolher")
    rt["A1"] = f"Retenções feitas pela {nome_emp} como tomadora — recolher/declarar"
    rt["A1"].font = titulo
    for j, t in enumerate(["Nº NFS-e", "Competência", "Prestador", "Valor do serviço", "ISS retido", "IRRF retido",
                           "PIS/COFINS/CSLL retidos", "Conta a crédito sugerida"], 1):
        cel = rt.cell(3, j, t)
        cel.font = Font(name="Arial", size=11, bold=True, color="FFFFFFFF"); cel.fill = cab_fundo
        cel.alignment = Alignment(wrap_text=True, horizontal="center")
    for col, w in zip("ABCDEFGH", (16, 12, 40, 14, 12, 12, 14, 60)):
        rt.column_dimensions[col].width = w
    lr = 4
    for lin, n, r in resultados:
        if r["confianca"] == "Excluir" or not (n["iss_retido"] or n["irrf_retido"] or n["csrf_retido"]):
            continue
        refs = [f"={aba}!A{lin}", f"={aba}!B{lin}", f"={aba}!D{lin}", f"={aba}!H{lin}",
                f"={aba}!I{lin}", f"={aba}!J{lin}", f"={aba}!K{lin}", r["retencoes"]]
        for j, v in enumerate(refs, start=1):
            cel = rt.cell(lr, j, v)
            cel.font = Font(name="Arial", size=10)
            if 4 <= j <= 7:
                cel.number_format = "#,##0.00;\\-#,##0.00;\\-"
        lr += 1
    if lr == 4:
        rt.cell(4, 1, "Nenhuma nota com retenção.").font = Font(name="Arial", size=10, italic=True)
    return wb, resultados


def empresa_das_notas(notas):
    """A empresa tomadora que mais aparece nos XMLs: é de quem é o plano que se deve usar."""
    from collections import Counter
    c = Counter((n.get("tomador") or "", n.get("tomador_nome") or "") for n in notas if n.get("tomador"))
    if not c:
        return "", ""
    (doc, nome), _ = c.most_common(1)[0]
    return doc, nome
